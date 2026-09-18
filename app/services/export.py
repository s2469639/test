"""기안서 PDF / Word 출력."""

import os
import tempfile

from docx import Document


def export_word(proposal_draft):
    doc = Document()
    doc.add_heading("참가 기안서", level=1)
    doc.add_paragraph(proposal_draft.content or "")

    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path


def export_pdf(proposal_draft):
    """TODO: weasyprint(HTML->PDF)로 렌더링. 시스템 라이브러리 설치가 번거로우면
    pdfkit(+wkhtmltopdf)로 대체 가능."""
    from weasyprint import HTML

    html = f"<h1>참가 기안서</h1><p>{proposal_draft.content or ''}</p>"
    fd, path = tempfile.mkstemp(suffix=".pdf")
    os.close(fd)
    HTML(string=html).write_pdf(path)
    return path
